"""Build the offline Jana Address SQLite gazetteer from KATO and OSM snapshots."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sqlite3
import sys
import tempfile
import time
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "backend") not in sys.path:
    sys.path.insert(0, str(ROOT / "backend"))

from jana.gazetteer import normalize, normalize_house
from jana.street_types import canonical_street_type


_STREET_PREFIXES = (
    "улица ",
    "ул. ",
    "ул ",
    "көше ",
    "көшесі ",
    "проспект ",
    "просп. ",
    "пр. ",
    "пр-т ",
    "даңғыл ",
    "даңғылы ",
    "данғыл ",
    "данғылы ",
    "переулок ",
    "пер. ",
    "пер ",
    "проезд ",
    "шоссе ",
    "шоссесі ",
    "тас жолы ",
    "набережная ",
    "наб. ",
    "наб ",
    "жағалау ",
    "жағалауы ",
    "бульвар ",
    "б-р ",
    "аллея ",
    "тупик ",
    "тұйық ",
    "тұйығы ",
    "тұйық көшесі ",
    "площадь ",
    "пл. ",
    "пл ",
    "алаң ",
    "алаңы ",
)
_STREET_SUFFIXES = (
    " көшесі",
    " даңғылы",
    " данғылы",
    " тұйық көшесі",
    " тас жолы",
    " шоссесі",
    " улица",
    " проспект",
    " переулок",
    " проезд",
    " набережная",
    " жағалауы",
    " бульвар",
    " аллея",
    " тупик",
    " тұйығы",
    " площадь",
    " алаңы",
)


def street_alias_forms(value: str) -> tuple[str, ...]:
    """Return source-preserving and type-marker-free exact lookup aliases."""

    source = " ".join(value.split()).strip(" ,")
    forms = [source]
    folded = source.casefold()
    for marker in sorted(_STREET_PREFIXES, key=len, reverse=True):
        if folded.startswith(marker):
            stripped = source[len(marker) :].strip(" ,")
            if stripped:
                forms.append(stripped)
            break
    for marker in sorted(_STREET_SUFFIXES, key=len, reverse=True):
        if folded.endswith(marker):
            stripped = source[: -len(marker)].strip(" ,")
            if stripped:
                forms.append(stripped)
            break
    return tuple(dict.fromkeys(forms))


def split_street_surface(value: str) -> tuple[str, str | None]:
    """Split a catalog street type marker without rewriting the source name."""

    source = " ".join(value.split()).strip(" ,")
    folded = source.casefold()
    for marker in sorted(_STREET_PREFIXES, key=len, reverse=True):
        if folded.startswith(marker):
            name = source[len(marker) :].strip(" ,")
            marker_key = marker.strip()
            return name or source, canonical_street_type(marker_key)
    for marker in sorted(_STREET_SUFFIXES, key=len, reverse=True):
        if folded.endswith(marker):
            name = source[: -len(marker)].strip(" ,")
            marker_key = marker.strip()
            return name or source, canonical_street_type(marker_key)
    return source, None


@dataclass(frozen=True)
class AdminRecord:
    field: str
    canonical: str
    aliases: tuple[str, ...] = ()
    language: str = "ru"
    kato_code: str | None = None
    province: str | None = None
    city: str | None = None
    district: str | None = None
    living_area: str | None = None


@dataclass(frozen=True)
class StreetRecord:
    source_id: str
    name: str
    aliases: tuple[str, ...] = ()
    street_type: str | None = None
    city: str | None = None
    province: str | None = None
    living_area: str | None = None
    districts: tuple[str, ...] = ()
    geometry_wkb: bytes | None = None


@dataclass(frozen=True)
class AddressPointRecord:
    source_id: str
    street_source_id: str
    house: str
    district: str | None = None
    living_area: str | None = None
    longitude: float | None = None
    latitude: float | None = None


@dataclass(frozen=True)
class SourceRecord:
    source_name: str
    url: str
    filename: str
    sha256: str
    fetched_at: str
    version: str
    license_name: str
    license_url: str


_LOCALITY_PREFIX = re.compile(
    r"^(?:г\.|с\.|п\.|пос\.|ст\.|аул\s+|село\s+|город\s+)", re.IGNORECASE
)
_ADMIN_CONTAINER = re.compile(r"\b[Гг]\.\s*[Аа]\.\s*$")
_LOCAL_PLACES = {
    "city", "town", "village", "hamlet", "isolated_dwelling", "suburb",
    "quarter", "neighbourhood", "locality", "allotments",
}
_NONLOCAL_PLACES = {"country", "continent", "state", "province", "region"}


def _strip_locality_marker(value: str) -> str:
    return _LOCALITY_PREFIX.sub("", " ".join(value.split())).strip(" .,")


def parse_kato_workbook(path: str | Path) -> list[AdminRecord]:
    """Parse the official KATO workbook into Russian canonical admin records."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    columns = {name: index for index, name in enumerate(header)}
    required = {"te", "kaz_name", "rus_name"}
    if not required.issubset(columns):
        raise ValueError(f"KATO workbook is missing columns: {sorted(required - columns.keys())}")

    records: list[AdminRecord] = []
    current_province: str | None = None
    current_district: str | None = None
    for row in rows:
        code = str(row[columns["te"]] or "").strip()
        russian = " ".join(str(row[columns["rus_name"]] or "").split())
        kazakh = " ".join(str(row[columns["kaz_name"]] or "").split())
        if not code or not russian:
            continue
        aliases = tuple(dict.fromkeys(value for value in (russian, kazakh) if value))
        is_region = len(code) == 9 and code[2:] == "0" * 7
        if is_region:
            canonical = _strip_locality_marker(russian)
            current_province = canonical
            current_district = None
            records.append(AdminRecord(
                field="province",
                canonical=canonical,
                aliases=aliases,
                kato_code=code,
                province=canonical,
            ))
            if _LOCALITY_PREFIX.match(russian):
                records.append(AdminRecord(
                    field="city",
                    canonical=canonical,
                    aliases=aliases,
                    kato_code=code,
                    province=canonical,
                    city=canonical,
                ))
            continue

        if _ADMIN_CONTAINER.search(russian):
            if len(code) == 9 and code[3:] == "0" * 6:
                current_district = None
            continue
        if re.search(r"\bрайон\b", russian, re.IGNORECASE):
            current_district = russian.strip(" .,")
            records.append(AdminRecord(
                field="district",
                canonical=current_district,
                aliases=aliases,
                kato_code=code,
                province=current_province,
                district=current_district,
            ))
            continue
        if _LOCALITY_PREFIX.match(russian):
            canonical = _strip_locality_marker(russian)
            records.append(AdminRecord(
                field="city",
                canonical=canonical,
                aliases=tuple(dict.fromkeys((russian, kazakh, canonical))),
                kato_code=code,
                province=current_province,
                city=canonical,
                district=current_district,
            ))
            continue
        if re.search(r"(?:[сСаАпП]\.\s*[оОаА]\.)\s*$", russian):
            records.append(AdminRecord(
                field="living_area",
                canonical=russian.strip(" .,"),
                aliases=aliases,
                kato_code=code,
                province=current_province,
                district=current_district,
                living_area=russian.strip(" .,"),
            ))
    workbook.close()
    return records


def _admin_identity(record: AdminRecord) -> tuple[str | None, ...]:
    """Names/aliases identify an entity only together with its role and parents."""

    return (record.field, record.kato_code, record.canonical, record.province,
            record.city, record.district, record.living_area)


def _is_coextensive_province_city(records: Sequence[AdminRecord]) -> bool:
    """Recognize only the two roles emitted for a top-level KATO locality."""

    if len(records) != 2 or {record.field for record in records} != {"province", "city"}:
        return False
    by_field = {record.field: record for record in records}
    province, city = by_field["province"], by_field["city"]
    return bool(
        province.kato_code
        and re.fullmatch(r"[0-9]{2}0{7}", province.kato_code)
        and province.kato_code == city.kato_code
        and province.canonical.strip()
        and province.canonical == city.canonical == province.province == city.province == city.city
        and province.city is None
        and province.district is None and city.district is None
        and province.living_area is None and city.living_area is None
    )


def _canonical_admin_index(admins: Sequence[AdminRecord]) -> dict[str, tuple[AdminRecord, ...]]:
    grouped: dict[str, dict[tuple[str | None, ...], AdminRecord]] = defaultdict(dict)
    for record in admins:
        for alias in dict.fromkeys((record.canonical, *record.aliases)):
            forms = [normalize(alias)]
            if record.field == "district":
                stripped = re.sub(
                    r"^(?:район|р\s*н|аудан|ауданы|ауд)\s+(?:(?:им|имени)\s+)?",
                    "",
                    forms[0],
                )
                stripped = re.sub(
                    r"\s+(?:(?:атын|атындағы)\s+)?(?:район|аудан|ауданы|ауд)$",
                    "",
                    stripped,
                )
                if stripped:
                    forms.append(stripped)
            for normalized in dict.fromkeys(forms):
                if normalized:
                    grouped[normalized][_admin_identity(record)] = record
    return {alias: tuple(records.values()) for alias, records in grouped.items()}


class PreparedAdminIndex:
    """Scalar STRtree lookup with reusable, unsimplified administrative polygons.

    STRtree's predicate is evaluated as query-geometry versus indexed-geometry.
    Reversing within to contains allows GEOS to reuse the *static polygon's*
    prepared index for thousands of points. Intersects is symmetric. Invalid
    inputs deliberately retain the original GEOS path, without repair or loss.
    """

    def __init__(self, geometries: Sequence[Any]) -> None:
        from shapely import prepare
        from shapely.strtree import STRtree

        self.geometries = tuple(geometries)
        self.valid = tuple(geometry.is_valid for geometry in self.geometries)
        self.tree = STRtree(self.geometries)
        for geometry, valid in zip(self.geometries, self.valid):
            if valid:
                prepare(geometry)

    def query(self, geometry: Any, predicate: str) -> Any:
        if predicate not in {"within", "intersects"}:
            return self.tree.query(geometry, predicate=predicate)
        indexes = self.tree.query(geometry)
        if len(indexes) == 0:
            return indexes
        if not geometry.is_valid or any(not self.valid[int(index)] for index in indexes):
            return self.tree.query(geometry, predicate=predicate)
        if predicate == "within":
            matches = [self.geometries[int(index)].contains(geometry) for index in indexes]
        else:
            matches = [self.geometries[int(index)].intersects(geometry) for index in indexes]
        return indexes[matches]


class BuildProgress:
    """Small JSON-lines diagnostics; counts/timings only, never raw addresses."""

    def __init__(self) -> None:
        self.started = time.perf_counter()
        self.cpu_started = time.process_time()

    def emit(self, phase: str, event: str, **counts: int) -> None:
        print(json.dumps({
            "phase": phase,
            "event": event,
            "elapsed_seconds": round(time.perf_counter() - self.started, 3),
            "cpu_seconds": round(time.process_time() - self.cpu_started, 3),
            **counts,
        }), file=sys.stderr, flush=True)


def extract_osm_records(
    path: str | Path,
    admins: Sequence[AdminRecord],
) -> tuple[list[StreetRecord], list[AddressPointRecord], dict[str, int]]:
    """Stream an OSM snapshot and spatially join named roads/address points."""

    import osmium
    from shapely import wkb
    from shapely.geometry import Point, Polygon

    progress = BuildProgress()
    admin_index = _canonical_admin_index(admins)
    code_groups: dict[str, dict[tuple[str | None, ...], AdminRecord]] = defaultdict(dict)
    for record in admins:
        if record.kato_code:
            code_groups[record.kato_code][_admin_identity(record)] = record
    admin_codes = {
        code: tuple(records.values()) for code, records in code_groups.items()
    }
    shapes: dict[str, list[tuple[Any, AdminRecord]]] = defaultdict(list)
    place_nodes: list[tuple[Point, AdminRecord]] = []
    factory = osmium.geom.WKBFactory()

    def resolve_admin(tags: Any, allowed_fields: set[str]) -> tuple[AdminRecord, ...]:
        source_code = next(
            (
                tags.get(key)
                for key in ("kato", "ref:kato", "KATO")
                if tags.get(key) and re.fullmatch(r"[0-9]{9}", tags.get(key))
            ),
            None,
        )
        if source_code:
            code_matches = tuple(
                item for item in admin_codes.get(source_code, ())
                if item.field in allowed_fields
            )
            if len(code_matches) == 1 or _is_coextensive_province_city(code_matches):
                return code_matches

        parent_provinces: dict[tuple[str | None, ...], AdminRecord] = {}
        for key in ("addr:region", "addr:province", "is_in:state", "is_in:province"):
            value = tags.get(key)
            if not value:
                continue
            for item in admin_index.get(normalize(value), ()):
                if item.field == "province":
                    parent_provinces[_admin_identity(item)] = item

        for key in ("name:ru", "official_name:ru", "name", "name:kk", "official_name"):
            value = tags.get(key)
            if not value:
                continue
            matches = tuple(item for item in admin_index.get(normalize(value), ()) if item.field in allowed_fields)
            if len(matches) > 1 and len(parent_provinces) == 1:
                parent = next(iter(parent_provinces.values())).canonical
                matches = tuple(item for item in matches if item.province == parent)
            if len(matches) > 1 and source_code:
                prefix_lengths = {
                    item: len(os.path.commonprefix((source_code, item.kato_code or "")))
                    for item in matches
                }
                best = max(prefix_lengths.values(), default=0)
                if best >= 3:
                    scoped = tuple(
                        item for item in matches if prefix_lengths[item] == best
                    )
                    if len(scoped) == 1:
                        matches = scoped
            if len(matches) == 1 or _is_coextensive_province_city(matches):
                return matches
        return ()

    class BoundaryHandler(osmium.SimpleHandler):
        nodes_seen = 0
        areas_seen = 0

        def area(self, area: Any) -> None:
            self.areas_seen += 1
            if self.areas_seen % 10000 == 0:
                progress.emit("boundaries", "progress", nodes_seen=self.nodes_seen,
                              areas_seen=self.areas_seen,
                              admin_shapes=sum(len(entries) for entries in shapes.values()))
            tags = area.tags
            if tags.get("boundary") != "administrative" and not tags.get("place"):
                return
            place = tags.get("place")
            if place in {"country", "continent"}:
                return
            level_text = tags.get("admin_level")
            if level_text is None:
                # A name alone does not establish an administrative role.
                if place not in _LOCAL_PLACES:
                    return
                level = None
            else:
                if not re.fullmatch(r"[0-9]+", level_text):
                    return
                level = int(level_text)
                if level < 4:
                    return
            if level in {4, 5}:
                allowed = {"province", "city"}
            elif level in {6, 7}:
                allowed = {"district", "city"}
            else:
                allowed = {"city", "living_area", "district"}
            if place in _NONLOCAL_PLACES:
                allowed -= {"city", "living_area"}
            records = resolve_admin(tags, allowed)
            if not records:
                return
            try:
                geometry = wkb.loads(factory.create_multipolygon(area), hex=True)
            except (RuntimeError, ValueError):
                return
            if not geometry.is_empty:
                for record in records:
                    shapes[record.field].append((geometry, record))

        def node(self, node: Any) -> None:
            self.nodes_seen += 1
            if self.nodes_seen % 100000 == 0:
                progress.emit("boundaries", "progress", nodes_seen=self.nodes_seen,
                              areas_seen=self.areas_seen)
            place = node.tags.get("place")
            if not place or place in _NONLOCAL_PLACES or not node.location.valid():
                return
            for record in resolve_admin(node.tags, {"city", "living_area"}):
                place_nodes.append((Point(node.location.lon, node.location.lat), record))

    boundary_handler = BoundaryHandler()
    progress.emit("boundaries", "start")
    boundary_handler.apply_file(str(path), locations=True, idx="flex_mem")
    progress.emit("boundaries", "done", nodes_seen=boundary_handler.nodes_seen,
                  areas_seen=boundary_handler.areas_seen,
                  admin_shapes=sum(len(entries) for entries in shapes.values()))

    progress.emit("spatial_index", "start")
    trees: dict[str, tuple[Any, list[AdminRecord]]] = {}
    for field, entries in shapes.items():
        geometries = [entry[0] for entry in entries]
        trees[field] = (PreparedAdminIndex(geometries), [entry[1] for entry in entries])
    progress.emit("spatial_index", "done", fields=len(trees))

    def spatial_records(geometry: Any, field: str, predicate: str) -> tuple[AdminRecord, ...]:
        if field not in trees:
            return ()
        tree, records = trees[field]
        indexes = tree.query(geometry, predicate=predicate)
        distinct: dict[tuple[str | None, ...], AdminRecord] = {}
        for index in indexes:
            record = records[int(index)]
            distinct[_admin_identity(record)] = record
        return tuple(distinct.values())

    def compatible_province(record: AdminRecord, provinces: tuple[AdminRecord, ...]) -> bool:
        return len(provinces) <= 1 and (
            not provinces or not record.province or record.province == provinces[0].canonical
        )

    raw_roads: list[dict[str, Any]] = []
    raw_points: list[dict[str, Any]] = []

    class FeatureHandler(osmium.SimpleHandler):
        nodes_seen = 0
        ways_seen = 0

        def way(self, way: Any) -> None:
            self.ways_seen += 1
            if self.ways_seen % 10000 == 0:
                progress.emit("features", "progress", nodes_seen=self.nodes_seen,
                              ways_seen=self.ways_seen, named_roads=len(raw_roads),
                              address_points=len(raw_points))
            house = way.tags.get("addr:housenumber")
            address_street = way.tags.get("addr:street")
            if house and address_street:
                try:
                    coordinates = [
                        (node.lon, node.lat) for node in way.nodes if node.location.valid()
                    ]
                    if len(coordinates) >= 4 and coordinates[0] == coordinates[-1]:
                        address_point = Polygon(coordinates).representative_point()
                    elif coordinates:
                        address_point = Point(coordinates[0])
                    else:
                        address_point = None
                except (RuntimeError, ValueError):
                    address_point = None
                if address_point is not None:
                    raw_points.append({
                        "source_id": f"way/{way.id}",
                        "house": house,
                        "street": address_street,
                        "city": way.tags.get("addr:city"),
                        "point": address_point,
                    })
            if not way.tags.get("highway"):
                return
            surfaces = [
                way.tags.get(key)
                for key in ("name:ru", "name", "name:kk", "official_name", "alt_name")
                if way.tags.get(key)
            ]
            if not surfaces:
                return
            try:
                geometry = wkb.loads(factory.create_linestring(way), hex=True)
            except (RuntimeError, ValueError):
                return
            if geometry.is_empty:
                return
            canonical_surface = way.tags.get("name:ru") or way.tags.get("name") or surfaces[0]
            canonical, street_type = split_street_surface(canonical_surface)
            aliases = tuple(dict.fromkeys(
                form for surface in surfaces for form in street_alias_forms(surface)
            ))
            city_records = spatial_records(geometry.representative_point(), "city", "within")
            province_records = spatial_records(geometry.representative_point(), "province", "within")
            city_records = tuple(item for item in city_records if compatible_province(item, province_records))
            districts = spatial_records(geometry, "district", "intersects")
            living_areas = spatial_records(geometry.representative_point(), "living_area", "within")
            raw_roads.append({
                "source_id": f"way/{way.id}",
                "name": canonical,
                "aliases": aliases,
                "street_type": street_type,
                "city": city_records[0].canonical if len(city_records) == 1 else None,
                "city_record": city_records[0] if len(city_records) == 1 else None,
                "province": province_records[0].canonical if len(province_records) == 1 else None,
                "province_records": province_records,
                "living_area": living_areas[0].canonical if len(living_areas) == 1 else None,
                "districts": {item.canonical for item in districts},
                "geometry": geometry,
            })

        def node(self, node: Any) -> None:
            self.nodes_seen += 1
            if self.nodes_seen % 100000 == 0:
                progress.emit("features", "progress", nodes_seen=self.nodes_seen,
                              ways_seen=self.ways_seen, address_points=len(raw_points))
            house = node.tags.get("addr:housenumber")
            street = node.tags.get("addr:street")
            if not house or not street or not node.location.valid():
                return
            raw_points.append({
                "source_id": f"node/{node.id}",
                "house": house,
                "street": street,
                "city": node.tags.get("addr:city"),
                "point": Point(node.location.lon, node.location.lat),
            })

    feature_handler = FeatureHandler()
    progress.emit("features", "start")
    feature_handler.apply_file(str(path), locations=True, idx="flex_mem")
    progress.emit("features", "done", nodes_seen=feature_handler.nodes_seen,
                  ways_seen=feature_handler.ways_seen, named_roads=len(raw_roads),
                  address_points=len(raw_points))

    progress.emit("point_join", "start", address_points=len(raw_points))
    road_aliases: dict[str, list[int]] = defaultdict(list)
    for index, road in enumerate(raw_roads):
        for alias in road["aliases"]:
            normalized = normalize(alias)
            if normalized:
                road_aliases[normalized].append(index)

    point_records: list[AddressPointRecord] = []
    unresolved_points = 0
    road_cities: dict[int, set[str]] = defaultdict(set)
    road_provinces: dict[int, dict[tuple[str | None, ...], AdminRecord]] = defaultdict(dict)
    derived_roads: dict[tuple[tuple[str | None, ...], str], dict[str, Any]] = {}

    def add_derived_address_street(
        raw: dict[str, Any],
        city: AdminRecord,
        provinces: tuple[AdminRecord, ...],
        districts: tuple[AdminRecord, ...],
        living_areas: tuple[AdminRecord, ...],
    ) -> None:
        """Keep an explicit addr:street even when OSM has no matching highway."""

        canonical, street_type = split_street_surface(raw["street"])
        normalized_name = normalize(canonical)
        key = (_admin_identity(city), normalized_name)
        if key not in derived_roads:
            digest = hashlib.sha256(
                "\0".join(str(value or "") for value in (*_admin_identity(city), normalized_name)).encode("utf-8")
            ).hexdigest()[:24]
            derived_roads[key] = {
                "source_id": f"derived/addr-street/{digest}",
                "name": canonical,
                "aliases": set(street_alias_forms(raw["street"])),
                "street_type": street_type,
                "city": city.canonical,
                "province": provinces[0].canonical if len(provinces) == 1 else city.province,
                "districts": set(),
                "living_areas": set(),
            }
        road = derived_roads[key]
        road["aliases"].update(street_alias_forms(raw["street"]))
        if len(districts) == 1:
            road["districts"].add(districts[0].canonical)
        if len(living_areas) == 1:
            road["living_areas"].add(living_areas[0].canonical)
        point_records.append(AddressPointRecord(
            source_id=raw["source_id"],
            street_source_id=road["source_id"],
            house=raw["house"],
            district=districts[0].canonical if len(districts) == 1 else None,
            living_area=living_areas[0].canonical if len(living_areas) == 1 else None,
            longitude=raw["point"].x,
            latitude=raw["point"].y,
        ))

    for point_number, raw in enumerate(raw_points, 1):
        if point_number % 10000 == 0:
            progress.emit("point_join", "progress", points_processed=point_number - 1,
                          points_joined=len(point_records), unresolved_points=unresolved_points)
        point = raw["point"]
        city_matches = admin_index.get(normalize(raw["city"]), ()) if raw["city"] else ()
        cities = tuple(item for item in city_matches if item.field == "city")
        spatial_cities = spatial_records(point, "city", "within")
        if cities and spatial_cities:
            spatial_ids = {_admin_identity(item) for item in spatial_cities}
            cities = tuple(item for item in cities if _admin_identity(item) in spatial_ids)
        elif not cities:
            cities = spatial_cities
        provinces = spatial_records(point, "province", "within")
        cities = tuple(item for item in cities if compatible_province(item, provinces))
        districts = spatial_records(point, "district", "within")
        living_areas = spatial_records(point, "living_area", "within")
        candidates: set[int] = set()
        for alias in street_alias_forms(raw["street"]):
                candidates.update(road_aliases.get(normalize(alias), ()))
        if len(cities) != 1:
            unresolved_points += 1
            continue
        candidates = {
            index for index in candidates
            if raw_roads[index]["city_record"] is not None
            and _admin_identity(raw_roads[index]["city_record"]) == _admin_identity(cities[0])
            and compatible_province(cities[0], raw_roads[index]["province_records"])
            and (not provinces or not raw_roads[index]["province_records"]
                 or _admin_identity(provinces[0]) == _admin_identity(raw_roads[index]["province_records"][0]))
        }
        spatial_city_ids = {_admin_identity(item) for item in spatial_cities}
        can_derive_street = (
            len(spatial_city_ids) == 1
            and _admin_identity(cities[0]) in spatial_city_ids
        )
        if not candidates:
            if can_derive_street:
                add_derived_address_street(
                    raw, cities[0], provinces, districts, living_areas
                )
            else:
                unresolved_points += 1
            continue
        road_index = min(candidates, key=lambda index: raw_roads[index]["geometry"].distance(point))
        if raw_roads[road_index]["geometry"].distance(point) > 0.05:
            if can_derive_street:
                add_derived_address_street(
                    raw, cities[0], provinces, districts, living_areas
                )
            else:
                unresolved_points += 1
            continue
        if len(cities) == 1:
            road_cities[road_index].add(cities[0].canonical)
        if len(provinces) == 1:
            road_provinces[road_index][_admin_identity(provinces[0])] = provinces[0]
        point_records.append(AddressPointRecord(
            source_id=raw["source_id"],
            street_source_id=raw_roads[road_index]["source_id"],
            house=raw["house"],
            district=districts[0].canonical if len(districts) == 1 else None,
            living_area=living_areas[0].canonical if len(living_areas) == 1 else None,
            longitude=point.x,
            latitude=point.y,
        ))

    progress.emit("point_join", "done", points_processed=len(raw_points),
                  points_joined=len(point_records), unresolved_points=unresolved_points)
    progress.emit("finalize", "start", named_roads=len(raw_roads))
    street_records: list[StreetRecord] = []
    for index, road in enumerate(raw_roads):
        city = road["city"]
        if city is None and len(road_cities[index]) == 1:
            city = next(iter(road_cities[index]))
        province = road["province"]
        if province is None and len(road_provinces[index]) == 1:
            province = next(iter(road_provinces[index].values())).canonical
        if province is None and not road_provinces[index] and road["city_record"] is not None:
            province = road["city_record"].province
        street_records.append(StreetRecord(
            source_id=road["source_id"],
            name=road["name"],
            aliases=road["aliases"],
            street_type=road["street_type"],
            city=city,
            province=province,
            living_area=road["living_area"],
            districts=tuple(sorted(road["districts"])),
            geometry_wkb=road["geometry"].wkb,
        ))
    for road in derived_roads.values():
        living_areas = tuple(sorted(road["living_areas"]))
        street_records.append(StreetRecord(
            source_id=road["source_id"],
            name=road["name"],
            aliases=tuple(sorted(road["aliases"])),
            street_type=road["street_type"],
            city=road["city"],
            province=road["province"],
            living_area=living_areas[0] if len(living_areas) == 1 else None,
            districts=tuple(sorted(road["districts"])),
            geometry_wkb=None,
        ))
    report = {
        "admin_shapes": sum(len(entries) for entries in shapes.values()),
        "place_nodes": len(place_nodes),
        "named_street_ways": len(raw_roads),
        "derived_address_streets": len(derived_roads),
        "address_points_seen": len(raw_points),
        "address_points_joined": len(point_records),
        "unresolved_address_points": unresolved_points,
    }
    progress.emit("finalize", "done", streets=len(street_records),
                  address_points=len(point_records))
    return street_records, point_records, report


SCHEMA = """
PRAGMA journal_mode=DELETE;
PRAGMA synchronous=FULL;
CREATE TABLE source_snapshots (
  source_name TEXT PRIMARY KEY,
  url TEXT NOT NULL,
  filename TEXT NOT NULL,
  sha256 TEXT NOT NULL,
  fetched_at TEXT NOT NULL,
  version TEXT NOT NULL,
  license_name TEXT NOT NULL,
  license_url TEXT NOT NULL
);
CREATE TABLE admin_records (
  id INTEGER PRIMARY KEY,
  field TEXT NOT NULL CHECK(field IN ('province','city','district','living_area')),
  canonical TEXT NOT NULL,
  language TEXT NOT NULL,
  kato_code TEXT,
  province TEXT,
  city TEXT,
  district TEXT,
  living_area TEXT
);
CREATE TABLE admin_aliases (
  admin_id INTEGER NOT NULL REFERENCES admin_records(id),
  field TEXT NOT NULL,
  canonical TEXT NOT NULL,
  alias TEXT NOT NULL,
  normalized TEXT NOT NULL,
  language TEXT NOT NULL,
  kato_code TEXT,
  province TEXT,
  city TEXT,
  district TEXT,
  living_area TEXT,
  UNIQUE(admin_id, normalized)
);
CREATE INDEX admin_aliases_normalized_idx ON admin_aliases(normalized, field);
CREATE TABLE streets (
  id INTEGER PRIMARY KEY,
  source_id TEXT NOT NULL UNIQUE,
  name TEXT NOT NULL,
  street_type TEXT,
  city TEXT,
  normalized_city TEXT,
  province TEXT,
  living_area TEXT,
  geometry_wkb BLOB
);
CREATE TABLE street_aliases (
  street_id INTEGER NOT NULL REFERENCES streets(id),
  alias TEXT NOT NULL,
  normalized TEXT NOT NULL,
  UNIQUE(street_id, normalized)
);
CREATE INDEX street_aliases_normalized_idx ON street_aliases(normalized);
CREATE TABLE street_districts (
  street_id INTEGER NOT NULL REFERENCES streets(id),
  district TEXT NOT NULL,
  PRIMARY KEY(street_id, district)
);
CREATE TABLE address_points (
  id INTEGER PRIMARY KEY,
  source_id TEXT NOT NULL UNIQUE,
  street_id INTEGER NOT NULL REFERENCES streets(id),
  house TEXT NOT NULL,
  normalized_house TEXT NOT NULL,
  district TEXT,
  living_area TEXT,
  longitude REAL,
  latitude REAL
);
CREATE INDEX address_points_street_house_idx ON address_points(street_id, normalized_house);
"""


def unique_normalized_aliases(aliases: Iterable[str]) -> list[tuple[str, str]]:
    """Keep the first surface form for each nonempty per-entity lookup key."""
    by_key: dict[str, str] = {}
    for alias in aliases:
        key = normalize(alias)
        if key:
            by_key.setdefault(key, alias)
    return [(alias, key) for key, alias in by_key.items()]


def create_database(
    output: str | Path,
    admins: Iterable[AdminRecord],
    streets: Iterable[StreetRecord],
    address_points: Iterable[AddressPointRecord],
    sources: Iterable[SourceRecord],
) -> None:
    """Create the complete database atomically at *output*."""

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(prefix=f".{output_path.name}.", dir=output_path.parent)
    os.close(handle)
    temp_path = Path(temp_name)
    connection: sqlite3.Connection | None = None
    try:
        connection = sqlite3.connect(temp_path)
        connection.executescript(SCHEMA)
        connection.executemany(
            "INSERT INTO source_snapshots VALUES (?,?,?,?,?,?,?,?)",
            [
                (
                    item.source_name,
                    item.url,
                    item.filename,
                    item.sha256,
                    item.fetched_at,
                    item.version,
                    item.license_name,
                    item.license_url,
                )
                for item in sources
            ],
        )
        for item in admins:
            cursor = connection.execute(
                """
                INSERT INTO admin_records
                    (field,canonical,language,kato_code,province,city,district,living_area)
                VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    item.field,
                    item.canonical,
                    item.language,
                    item.kato_code,
                    item.province,
                    item.city,
                    item.district,
                    item.living_area,
                ),
            )
            aliases = unique_normalized_aliases((item.canonical, *item.aliases))
            connection.executemany(
                """
                INSERT INTO admin_aliases
                    (admin_id,field,canonical,alias,normalized,language,kato_code,
                     province,city,district,living_area)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
                """,
                [
                    (
                        cursor.lastrowid,
                        item.field,
                        item.canonical,
                        alias,
                        normalized,
                        item.language,
                        item.kato_code,
                        item.province,
                        item.city,
                        item.district,
                        item.living_area,
                    )
                    for alias, normalized in aliases
                ],
            )
        street_ids: dict[str, int] = {}
        for item in streets:
            cursor = connection.execute(
                """
                INSERT INTO streets
                    (source_id,name,street_type,city,normalized_city,province,living_area,geometry_wkb)
                VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    item.source_id,
                    item.name,
                    item.street_type,
                    item.city,
                    normalize(item.city) if item.city else None,
                    item.province,
                    item.living_area,
                    item.geometry_wkb,
                ),
            )
            street_ids[item.source_id] = int(cursor.lastrowid)
            aliases = unique_normalized_aliases(
                form
                for alias in (item.name, *item.aliases)
                for form in street_alias_forms(alias)
            )
            connection.executemany(
                "INSERT INTO street_aliases(street_id,alias,normalized) VALUES (?,?,?)",
                [
                    (cursor.lastrowid, alias, normalized)
                    for alias, normalized in aliases
                ],
            )
            connection.executemany(
                "INSERT INTO street_districts(street_id,district) VALUES (?,?)",
                [(cursor.lastrowid, district) for district in dict.fromkeys(item.districts)],
            )
        for item in address_points:
            street_id = street_ids.get(item.street_source_id)
            if street_id is None:
                continue
            connection.execute(
                """
                INSERT INTO address_points
                    (source_id,street_id,house,normalized_house,district,living_area,longitude,latitude)
                VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    item.source_id,
                    street_id,
                    item.house,
                    normalize_house(item.house),
                    item.district,
                    item.living_area,
                    item.longitude,
                    item.latitude,
                ),
            )
        connection.execute("ANALYZE")
        connection.commit()
        connection.close()
        connection = None
        os.replace(temp_path, output_path)
    finally:
        if connection is not None:
            connection.close()
        if temp_path.exists():
            temp_path.unlink()


def _load_json_records(path: Path) -> tuple[list[AdminRecord], list[StreetRecord], list[AddressPointRecord]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return (
        [AdminRecord(**row) for row in payload.get("admins", [])],
        [StreetRecord(**row) for row in payload.get("streets", [])],
        [AddressPointRecord(**row) for row in payload.get("address_points", [])],
    )


def load_verified_contexts(path: str | Path) -> tuple[list[StreetRecord], list[AddressPointRecord]]:
    """Convert source-attributed exact address contexts into gazetteer rows.

    A district is attached to an exact house point, never to the whole street.
    This prevents a verified building from making up district membership for
    other houses on a street that may cross an administrative boundary.
    """

    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if payload.get("schema_version") != 1 or not isinstance(payload.get("contexts"), list):
        raise ValueError("verified address contexts must use schema_version 1")
    streets: list[StreetRecord] = []
    points: list[AddressPointRecord] = []
    for row in payload["contexts"]:
        required = ("source_id", "city", "district", "street_name", "street_type", "houses")
        missing = [field for field in required if not row.get(field)]
        if missing:
            raise ValueError(f"verified address context is missing fields: {missing}")
        source_id = str(row["source_id"])
        streets.append(StreetRecord(
            source_id=source_id,
            name=str(row["street_name"]),
            aliases=tuple(str(value) for value in row.get("aliases", ())),
            street_type=str(row["street_type"]),
            city=str(row["city"]),
            province=str(row["province"]) if row.get("province") else None,
        ))
        for house in row["houses"]:
            if not house.get("source_id") or not house.get("house"):
                raise ValueError("verified house requires source_id and house")
            points.append(AddressPointRecord(
                source_id=str(house["source_id"]),
                street_source_id=source_id,
                house=str(house["house"]),
                district=str(row["district"]),
                longitude=float(house["longitude"]) if house.get("longitude") is not None else None,
                latitude=float(house["latitude"]) if house.get("latitude") is not None else None,
            ))
    return streets, points


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--fixture-json", type=Path)
    inputs.add_argument("--kato", type=Path)
    parser.add_argument("--osm", type=Path)
    parser.add_argument("--verified-contexts", type=Path)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args(argv)
    started = time.perf_counter()
    progress = BuildProgress()
    osm_report: dict[str, int] = {}
    progress.emit("input", "start")
    if args.fixture_json:
        admins, streets, points = _load_json_records(args.fixture_json)
        progress.emit("input", "done", admin_records=len(admins))
    else:
        if args.osm is None:
            parser.error("--osm is required with --kato")
        admins = parse_kato_workbook(args.kato)
        progress.emit("input", "done", admin_records=len(admins))
        streets, points, osm_report = extract_osm_records(args.osm, admins)
    if args.verified_contexts:
        verified_streets, verified_points = load_verified_contexts(args.verified_contexts)
        streets.extend(verified_streets)
        points.extend(verified_points)
        osm_report["verified_context_streets"] = len(verified_streets)
        osm_report["verified_context_points"] = len(verified_points)
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    sources = [SourceRecord(**row) for row in manifest["sources"]]
    progress.emit("database", "start", admin_records=len(admins), streets=len(streets),
                  address_points=len(points))
    create_database(args.output, admins, streets, points, sources)
    progress.emit("database", "done", database_bytes=args.output.stat().st_size)
    elapsed = time.perf_counter() - started
    report = {
        "database": str(args.output),
        "build_seconds": round(elapsed, 3),
        "admin_records": len(admins),
        "streets": len(streets),
        "address_points": len(points),
        **osm_report,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
